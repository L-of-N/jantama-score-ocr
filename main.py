# -*- coding: utf-8 -*-
import os
import sys
import ssl
import hashlib

import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog



def get_app_dir():
    """Return the folder that should contain images/output/templates."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)

    return os.path.dirname(os.path.abspath(__file__))


def select_run_settings_gui(default_image_dir, default_output_dir, default_template_dir):
    """Show a small launcher so non-technical users can run the OCR app."""
    if "--no-gui" in sys.argv:
        return default_image_dir, default_output_dir, default_template_dir

    selected = {"run": False}

    root = tk.Tk()
    root.title("雀魂スコアOCR")
    root.geometry("620x260")
    root.resizable(False, False)
    root.attributes("-topmost", True)

    image_var = tk.StringVar(value=default_image_dir)
    output_var = tk.StringVar(value=default_output_dir)
    template_var = tk.StringVar(value=default_template_dir)

    def browse_folder(var):
        folder = filedialog.askdirectory(initialdir=var.get() or get_app_dir())
        if folder:
            var.set(folder)

    def add_row(row, label_text, var):
        tk.Label(root, text=label_text, anchor="w", width=16).grid(
            row=row,
            column=0,
            padx=12,
            pady=8,
            sticky="w"
        )
        tk.Entry(root, textvariable=var, width=58).grid(
            row=row,
            column=1,
            padx=4,
            pady=8,
            sticky="we"
        )
        tk.Button(
            root,
            text="選択",
            command=lambda: browse_folder(var),
            width=8
        ).grid(row=row, column=2, padx=8, pady=8)

    tk.Label(
        root,
        text="画像フォルダを選んで実行してください。結果は出力フォルダに保存されます。",
        anchor="w"
    ).grid(row=0, column=0, columnspan=3, padx=12, pady=(14, 8), sticky="w")

    add_row(1, "画像フォルダ", image_var)
    add_row(2, "出力フォルダ", output_var)
    add_row(3, "テンプレート", template_var)

    def run():
        image_dir = image_var.get().strip()
        output_dir = output_var.get().strip()
        template_dir = template_var.get().strip()

        if not os.path.isdir(image_dir):
            messagebox.showerror("確認", "画像フォルダが見つかりません。")
            return

        if not os.path.isdir(template_dir):
            messagebox.showerror("確認", "テンプレートフォルダが見つかりません。")
            return

        missing_templates = [
            f"{rank}.png"
            for rank in range(1, 5)
            if not os.path.exists(os.path.join(template_dir, f"{rank}.png"))
        ]

        if missing_templates:
            messagebox.showerror(
                "確認",
                "テンプレート画像が不足しています: " + ", ".join(missing_templates)
            )
            return

        os.makedirs(output_dir, exist_ok=True)
        selected["run"] = True
        selected["image_dir"] = image_dir
        selected["output_dir"] = output_dir
        selected["template_dir"] = template_dir
        root.destroy()

    def cancel():
        root.destroy()

    button_frame = tk.Frame(root)
    button_frame.grid(row=4, column=0, columnspan=3, pady=16)

    tk.Button(button_frame, text="実行", command=run, width=14).pack(side="left", padx=8)
    tk.Button(button_frame, text="キャンセル", command=cancel, width=14).pack(side="left", padx=8)

    root.protocol("WM_DELETE_WINDOW", cancel)
    root.mainloop()

    if not selected["run"]:
        sys.exit(0)

    return selected["image_dir"], selected["output_dir"], selected["template_dir"]

def bind_vertical_mousewheel(widget, canvas):
    """Scroll a canvas with the mouse wheel while the pointer is over it."""
    def on_mousewheel(event):
        if event.num == 4:
            canvas.yview_scroll(-1, "units")
        elif event.num == 5:
            canvas.yview_scroll(1, "units")
        else:
            canvas.yview_scroll(int(-event.delta / 120), "units")

    def bind_events(_event):
        canvas.bind_all("<MouseWheel>", on_mousewheel)
        canvas.bind_all("<Button-4>", on_mousewheel)
        canvas.bind_all("<Button-5>", on_mousewheel)

    def unbind_events(_event):
        canvas.unbind_all("<MouseWheel>")
        canvas.unbind_all("<Button-4>")
        canvas.unbind_all("<Button-5>")

    widget.bind("<Enter>", bind_events)
    widget.bind("<Leave>", unbind_events)


def select_players_gui(player_names):
    selected_players = []

    root = tk.Tk()
    root.title("出力するプレイヤーを選択")
    root.geometry("460x650")
    root.minsize(360, 360)

    label = tk.Label(
        root,
        text="個別シートを作成するプレイヤーを選択してください"
    )
    label.pack(pady=10)

    vars_dict = {}

    outer_frame = tk.Frame(root)
    outer_frame.pack(fill="both", expand=True, padx=10)

    canvas = tk.Canvas(outer_frame, highlightthickness=0)
    scrollbar = tk.Scrollbar(
        outer_frame,
        orient="vertical",
        command=canvas.yview
    )
    frame = tk.Frame(canvas)
    frame.bind(
        "<Configure>",
        lambda _event: canvas.configure(scrollregion=canvas.bbox("all"))
    )
    frame_window = canvas.create_window((0, 0), window=frame, anchor="nw")
    canvas.bind(
        "<Configure>",
        lambda event: canvas.itemconfigure(frame_window, width=event.width)
    )
    canvas.configure(yscrollcommand=scrollbar.set)
    canvas.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")
    bind_vertical_mousewheel(canvas, canvas)

    for name in player_names:
        var = tk.BooleanVar(value=True)
        chk = tk.Checkbutton(
            frame,
            text=name,
            variable=var,
            anchor="w"
        )
        chk.pack(fill="x", padx=12, pady=1)

        vars_dict[name] = var

    def select_all():
        for var in vars_dict.values():
            var.set(True)

    def clear_all():
        for var in vars_dict.values():
            var.set(False)

    def execute():
        for name, var in vars_dict.items():
            if var.get():
                selected_players.append(name)

        root.destroy()

    btn_frame = tk.Frame(root)
    btn_frame.pack(pady=10)

    tk.Button(
        btn_frame,
        text="全選択",
        command=select_all
    ).pack(side="left", padx=5)

    tk.Button(
        btn_frame,
        text="全解除",
        command=clear_all
    ).pack(side="left", padx=5)

    tk.Button(
        root,
        text="この内容で作成",
        command=execute
    ).pack(pady=10)

    root.mainloop()

    return selected_players


def ask_name_manual_gui(filename, rank):
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    value = simpledialog.askstring(
        "名前OCR失敗",
        f"{filename} の {rank}位の名前が読めませんでした。\n名前を入力してください。",
        parent=root
    )

    root.destroy()

    if value is None:
        return ""

    return clean_name_text(value)

ssl._create_default_https_context = ssl._create_unverified_context
os.environ["OMP_NUM_THREADS"] = "1"
def configure_console_encoding():
    """Keep Japanese console output readable on Windows UTF-8 consoles."""
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


configure_console_encoding()

import cv2
import numpy as np
import pytesseract
from PIL import Image, ImageTk
import re
import pandas as pd
import easyocr

from openpyxl.styles import Border, Side, Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

pytesseract.pytesseract.tesseract_cmd = os.environ.get(
    "TESSERACT_CMD",
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)

APP_DIR = get_app_dir()
IMAGE_DIR, OUTPUT_DIR, TEMPLATE_DIR = select_run_settings_gui(
    os.path.join(APP_DIR, "images"),
    os.path.join(APP_DIR, "output"),
    os.path.join(APP_DIR, "templates")
)
ALIAS_PATH = os.path.join(OUTPUT_DIR, "name_alias.csv")

os.makedirs(OUTPUT_DIR, exist_ok=True)

easy_reader = easyocr.Reader(
    ["ja", "en"],
    gpu=False
)

rows = []


def imread_unicode(path):
    """Read an image even when its Windows path contains Japanese text."""
    try:
        data = np.fromfile(path, dtype=np.uint8)
        return cv2.imdecode(data, cv2.IMREAD_COLOR)
    except Exception:
        return None


def imwrite_unicode(path, image):
    """Write an image even when its Windows path contains Japanese text."""
    extension = os.path.splitext(path)[1] or ".png"
    success, encoded = cv2.imencode(extension, image)

    if success:
        encoded.tofile(path)

    return success


def get_image_hash(image_path):
    with open(image_path, "rb") as f:
        return hashlib.md5(f.read()).hexdigest()


def clean_name_text(text):
    text = str(text).strip()
    text = text.replace(" ", "")
    text = text.replace("　", "")
    text = text.replace("－", "-")

    text = re.sub(
        r"[^ぁ-んァ-ン一-龯a-zA-Z0-9ー\-_@]",
        "",
        text
    )

    return text



def load_name_aliases():
    """name_alias.csv から OCR名 -> 正しい名前 の辞書を読み込む。"""
    if not os.path.exists(ALIAS_PATH):
        return {}

    try:
        alias_df = pd.read_csv(ALIAS_PATH, encoding="utf-8-sig")
    except Exception:
        return {}

    if "OCR名" not in alias_df.columns or "正しい名前" not in alias_df.columns:
        return {}

    alias_dict = {}

    for _, row in alias_df.iterrows():
        ocr_name = clean_name_text(row.get("OCR名", ""))
        correct_name = clean_name_text(row.get("正しい名前", ""))

        if ocr_name and correct_name:
            alias_dict[ocr_name] = correct_name

    return alias_dict


def save_name_aliases(alias_dict):
    """OCR名 -> 正しい名前 の辞書を name_alias.csv に保存する。"""
    if not alias_dict:
        return

    alias_rows = []

    for ocr_name, correct_name in sorted(alias_dict.items()):
        ocr_name = clean_name_text(ocr_name)
        correct_name = clean_name_text(correct_name)

        if ocr_name and correct_name and ocr_name != correct_name:
            alias_rows.append({
                "OCR名": ocr_name,
                "正しい名前": correct_name
            })

    alias_df = pd.DataFrame(alias_rows)

    alias_df.to_csv(
        ALIAS_PATH,
        index=False,
        encoding="utf-8-sig"
    )


def confirm_names_gui(player_names, name_sources=None):
    """
    OCRされた名前一覧を表示して、必要なものだけ修正する。
    入力欄が空欄なら、そのOCR名は正しいものとして扱う。
    戻り値: {OCR名: 修正後の名前}
    """
    player_names = [
        clean_name_text(name)
        for name in player_names
        if clean_name_text(name)
    ]

    player_names = sorted(set(player_names))

    if not player_names:
        return {}

    corrections = {}
    name_sources = name_sources or {}

    root = tk.Tk()
    root.title("名前OCR確認")
    root.geometry("1120x680")
    root.minsize(850, 420)
    root.attributes("-topmost", True)

    guide = tk.Label(
        root,
        text="OCR結果が違う名前だけ右側に入力してください。空欄ならそのまま採用します。",
        wraplength=1080,
        justify="left"
    )
    guide.pack(padx=10, pady=10, anchor="w")

    outer_frame = tk.Frame(root)
    outer_frame.pack(fill="both", expand=True, padx=10)

    canvas = tk.Canvas(outer_frame, highlightthickness=0)
    scrollbar = tk.Scrollbar(
        outer_frame,
        orient="vertical",
        command=canvas.yview
    )

    scroll_frame = tk.Frame(canvas)

    scroll_frame.bind(
        "<Configure>",
        lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
    )

    scroll_frame.columnconfigure(0, weight=1)
    scroll_frame.columnconfigure(1, weight=1)
    scroll_frame.columnconfigure(2, weight=2)
    frame_window = canvas.create_window((0, 0), window=scroll_frame, anchor="nw")
    canvas.bind(
        "<Configure>",
        lambda event: canvas.itemconfigure(frame_window, width=event.width)
    )
    canvas.configure(yscrollcommand=scrollbar.set)

    canvas.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")
    bind_vertical_mousewheel(canvas, canvas)

    tk.Label(
        scroll_frame,
        text="OCR結果",
        width=24,
        anchor="w"
    ).grid(row=0, column=0, padx=5, pady=5, sticky="w")

    tk.Label(
        scroll_frame,
        text="修正後（空欄ならそのまま）",
        width=30,
        anchor="w"
    ).grid(row=0, column=1, padx=5, pady=5, sticky="w")

    tk.Label(
        scroll_frame,
        text="元画像（最初の1件）",
        width=48,
        anchor="w"
    ).grid(row=0, column=2, padx=5, pady=5, sticky="w")

    entry_dict = {}

    def show_extracted_image(name):
        sources = name_sources.get(name, [])

        if not sources:
            messagebox.showinfo(
                "抽出画像なし",
                "表示できる抽出画像がありません。",
                parent=root
            )
            return

        source = sources[0]
        preview_path = source.get("preview_path", "")

        if not preview_path or not os.path.exists(preview_path):
            messagebox.showinfo(
                "抽出画像なし",
                "抽出画像が見つかりません。過去に取り込んだ画像では表示できない場合があります。",
                parent=root
            )
            return

        source_window = tk.Toplevel(root)
        source_window.title(f"抽出画像: {name}")
        source_window.transient(root)
        source_window.attributes("-topmost", True)

        tk.Label(
            source_window,
            text=f"{name} / {source.get('label', '')}",
            anchor="w"
        ).pack(fill="x", padx=10, pady=(10, 5))

        image = Image.open(preview_path)
        image.thumbnail((900, 500))
        photo = ImageTk.PhotoImage(image)
        image_label = tk.Label(source_window, image=photo)
        image_label.image = photo
        image_label.pack(padx=10, pady=5)

        tk.Button(
            source_window,
            text="閉じる",
            command=source_window.destroy,
            width=12
        ).pack(pady=10)

        source_window.lift()
        source_window.focus_force()

    for i, name in enumerate(player_names, start=1):
        tk.Label(
            scroll_frame,
            text=name,
            width=24,
            anchor="w"
        ).grid(row=i, column=0, padx=5, pady=3, sticky="we")

        entry = tk.Entry(scroll_frame, width=30)
        entry.grid(row=i, column=1, padx=5, pady=3, sticky="we")

        entry_dict[name] = entry

        sources = name_sources.get(name, [])
        preview = sources[0].get("label", "") if sources else ""

        source_frame = tk.Frame(scroll_frame)
        source_frame.grid(row=i, column=2, padx=5, pady=3, sticky="we")

        tk.Label(
            source_frame,
            text=preview,
            width=42,
            anchor="w"
        ).pack(side="left", fill="x", expand=True)

        tk.Button(
            source_frame,
            text="抽出画像を表示",
            command=lambda current_name=name: show_extracted_image(current_name),
            width=12
        ).pack(side="right", padx=(5, 0))

    def execute():
        for ocr_name, entry in entry_dict.items():
            fixed_name = clean_name_text(entry.get())

            if fixed_name and fixed_name != ocr_name:
                corrections[ocr_name] = fixed_name

        root.destroy()

    btn_frame = tk.Frame(root)
    btn_frame.pack(pady=10)

    tk.Button(
        btn_frame,
        text="この内容で確定",
        command=execute,
        width=18
    ).pack(side="left", padx=5)

    root.mainloop()

    return corrections


def apply_excel_styles(workbook):
    """summary_result.xlsx 全体に罫線・列幅・見出し装飾を付ける。"""
    thin = Side(border_style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAD3")
    stat_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    rank_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")

    percent_labels = {"トップ率", "ラス率", "飛び率", "勝率"}

    def set_header_style(cell):
        if cell.value is not None:
            cell.fill = header_fill
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def format_percent_cell(cell):
        """50 -> 50% 表示になるように、値を 0.5 に直して%書式を付ける。"""
        if cell.value is None or cell.value == "":
            return

        try:
            value_text = str(cell.value).replace("%", "").strip()
            value = float(value_text)
        except Exception:
            return

        # すでに 0.5 のような値ならそのまま、50 のような値なら 100で割る
        if abs(value) > 1:
            value = value / 100

        cell.value = value
        cell.number_format = "0%"

    for ws in workbook.worksheets:
        # 値が入っているセルだけ罫線・配置
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cell.border = border
                    cell.alignment = Alignment(vertical="center", wrap_text=True)

        # 通常シートの1行目ヘッダー
        for cell in ws[1]:
            set_header_style(cell)

        # ％列を整形（summary / vs_summary / P_シートなど）
        for row in ws.iter_rows():
            for cell in row:
                if cell.value in percent_labels:
                    # 横持ち表：ヘッダーの下を％表示
                    for target in ws.iter_rows(
                        min_row=cell.row + 1,
                        max_row=ws.max_row,
                        min_col=cell.column,
                        max_col=cell.column
                    ):
                        format_percent_cell(target[0])

                    # 戦績カード：A列に「トップ率」等、B列に値がある場合
                    if cell.column == 1:
                        format_percent_cell(ws.cell(row=cell.row, column=2))

        if ws.title.startswith("戦績_"):
            # 戦績シートだけ14行目を見出しとして色付け
            for col in range(1, 6):
                set_header_style(ws.cell(row=14, column=col))

            for col in range(7, 13):
                set_header_style(ws.cell(row=14, column=col))

            # カード部分の色分け
            for row in ws.iter_rows(min_row=3, max_row=11, min_col=1, max_col=2):
                for cell in row:
                    if cell.value is not None:
                        cell.fill = stat_fill

            for row in ws.iter_rows(min_row=3, max_row=7, min_col=4, max_col=5):
                for cell in row:
                    if cell.value is not None:
                        cell.fill = rank_fill

            # 見出しは太字
            ws["A1"].font = Font(bold=True, size=14)
            ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

            for cell in [ws["A3"], ws["B3"], ws["D3"], ws["E3"]]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # 数値を右寄せ、項目名を左寄せ
            for cell in ws["A"]:
                if cell.value is not None:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            for cell in ws["B"]:
                if cell.value is not None:
                    cell.alignment = Alignment(horizontal="right", vertical="center")

            # 戦績シートは、手動調整済みの summary_result.xlsx に合わせた固定幅にする
            fixed_widths = {
                "A": 15.625,  # 項目
                "B": 9.75,    # 値
                "C": 7.5,
                "D": 7.625,   # 順位
                "E": 5.75,    # 回数
                "F": 3,
                "G": 15,      # 画像ファイル
                "H": 5.75,    # 順位
                "I": 7.5,     # 素点
                "J": 7.75,    # 順位点
                "K": 10.75,   # 確認フラグ
                "L": 34.75,   # 確認理由
            }

            for col, width in fixed_widths.items():
                ws.column_dimensions[col].width = width

            # 行の高さを少しだけ広げる
            for row_idx in range(1, ws.max_row + 1):
                ws.row_dimensions[row_idx].height = 18

            ws.freeze_panes = "A14"

        else:
            # 通常シートは13行目に不要な色を付けない
            ws.freeze_panes = "A2"

            # 手動調整済みの summary_result.xlsx に合わせて、主要シートは固定幅にする
            fixed_widths_by_sheet = {
                "summary": {
                    "A": 15.625,  # 名前OCR
                    "B": 7.75,    # 対局数
                    "C": 8.875,   # 1位回数
                    "D": 13,      # 2位回数
                    "E": 13,      # 3位回数
                    "F": 13,      # 4位回数
                    "G": 8.25,    # トップ率
                    "H": 7.125,   # ラス率
                    "I": 9.625,   # 飛び回数
                    "J": 7.625,   # 飛び率
                    "K": 9.75,    # 平均順位
                    "L": 13,      # 素点合計
                    "M": 11.875,  # 順位点合計
                },
                "vs_summary": {
                    "A": 15.625,  # 名前OCR
                    "B": 13,      # 対戦相手
                    "C": 9.75,    # 同卓回数
                    "D": 7.5,     # 勝ち数
                    "E": 7.625,   # 負け数
                    "F": 5.75,    # 勝率
                },
                "vs_raw": {
                    "A": 15.625,  # 名前OCR
                    "B": 15.625,  # 対戦相手
                    "C": 35.125,  # 画像ハッシュ
                    "D": 9.75,    # 自分順位
                    "E": 13,      # 相手順位
                    "F": 5.5,     # 勝ち
                    "G": 5.625,   # 負け
                },
            }

            if ws.title.startswith("P_"):
                fixed_widths = {
                    "A": 15.625,  # 対戦相手
                    "B": 9.75,    # 同卓回数
                    "C": 7.5,     # 勝ち数
                    "D": 7.625,   # 負け数
                    "E": 5.75,    # 勝率
                }
            else:
                fixed_widths = fixed_widths_by_sheet.get(ws.title, {})

            # 固定幅がないシートだけ、従来通り内容からざっくり自動調整する
            if fixed_widths:
                for col, width in fixed_widths.items():
                    ws.column_dimensions[col].width = width
            else:
                for col_cells in ws.columns:
                    max_len = 0
                    col_letter = get_column_letter(col_cells[0].column)

                    for cell in col_cells:
                        if cell.value is None:
                            continue

                        value_len = len(str(cell.value))
                        if value_len > max_len:
                            max_len = value_len

                    ws.column_dimensions[col_letter].width = min(max(max_len + 2, 8), 28)

                # 名前・相手名系は最低15にする
                for row in ws.iter_rows(min_row=1, max_row=1):
                    for cell in row:
                        if cell.value in ["名前OCR", "対戦相手", "プレイヤー"]:
                            col_letter = get_column_letter(cell.column)
                            ws.column_dimensions[col_letter].width = max(
                                ws.column_dimensions[col_letter].width,
                                15
                            )

            for row_idx in range(1, ws.max_row + 1):
                ws.row_dimensions[row_idx].height = 18



def ocr_name_easy(name_img, save_name):
    temp_path = os.path.join(OUTPUT_DIR, save_name)

    # 名前OCRは前処理を強くしすぎると、雀魂の縁取り文字が潰れることがあるため、
    # いったん「拡大 + グレースケール」だけに戻しています。
    big = cv2.resize(name_img, None, fx=5, fy=5)
    gray = cv2.cvtColor(big, cv2.COLOR_BGR2GRAY)

    imwrite_unicode(temp_path, gray)

    print("名前EasyOCR開始:", save_name)

    result = easy_reader.readtext(
        gray,
        detail=1,
        paragraph=False
    )

    print("名前EasyOCR結果:", result)
    print("名前EasyOCR終了:", save_name)

    texts = []

    for item in result:
        try:
            text = item[1]
            score = item[2]

            # 低信頼でも拾えるように少し緩める。
            # 文字化けが増える場合は 0.15 や 0.20 に上げてください。
            if score >= 0.10:
                texts.append(text)

        except Exception:
            pass

    final_text = clean_name_text("".join(texts))

    print("名前EasyOCR:", final_text)

    return final_text


def ocr_digits_from_mask(mask_img):
    big = cv2.resize(mask_img, None, fx=5, fy=5)

    text = pytesseract.image_to_string(
        Image.fromarray(big),
        lang="eng",
        config="--psm 7 -c tessedit_char_whitelist=0123456789."
    ).strip()

    nums = re.findall(r"\d+(?:\.\d+)?", text)

    return text, nums[0] if nums else ""


def detect_minus_from_score_bin(score_bin):
    black = (score_bin < 128).astype("uint8") * 255

    num_labels, labels, stats, _ = cv2.connectedComponentsWithStats(
        black,
        8
    )

    components = []

    for i in range(1, num_labels):
        x = stats[i, cv2.CC_STAT_LEFT]
        y = stats[i, cv2.CC_STAT_TOP]
        w = stats[i, cv2.CC_STAT_WIDTH]
        h = stats[i, cv2.CC_STAT_HEIGHT]
        area = stats[i, cv2.CC_STAT_AREA]

        if area < 20:
            continue

        components.append((x, y, w, h, area))

    digit_candidates = [
        c for c in components
        if c[2] >= 25 and c[3] >= 60
    ]

    if not digit_candidates:
        return False

    first_digit_x = min(c[0] for c in digit_candidates)

    for x, y, w, h, area in components:
        aspect = w / h if h else 0

        is_minus_shape = (
            w >= 40 and
            8 <= h <= 45 and
            aspect >= 2.0 and
            x < first_digit_x
        )

        if is_minus_shape:
            return True

    return False


# =========================
# 既存master読み込み
# =========================

old_dfs = {}
registered_hashes = set()

for game_type in ["yonma", "sanma"]:
    master_path = os.path.join(
        OUTPUT_DIR,
        game_type,
        "master_result.xlsx"
    )

    if os.path.exists(master_path):
        old_df = pd.read_excel(master_path)
    elif game_type == "yonma":
        legacy_master_path = os.path.join(
            OUTPUT_DIR,
            "master_result.xlsx"
        )

        if os.path.exists(legacy_master_path):
            old_df = pd.read_excel(legacy_master_path)
            print("旧版の四麻履歴を引き継ぎます:", legacy_master_path)
        else:
            old_df = pd.DataFrame()
    else:
        old_df = pd.DataFrame()

    old_dfs[game_type] = old_df

    if "画像ハッシュ" in old_df.columns:
        registered_hashes.update(
            old_df["画像ハッシュ"].astype(str).dropna()
        )


def detect_game_type(img_cv):
    """Classify a result screen by whether a real fourth-place panel exists."""
    template = cv2.imread(os.path.join(TEMPLATE_DIR, "4.png"))

    if template is None:
        raise RuntimeError("4位テンプレートが読み込めません。")

    result = cv2.matchTemplate(
        img_cv,
        template,
        cv2.TM_CCOEFF_NORMED
    )
    _, max_val, _, _ = cv2.minMaxLoc(result)
    game_type = "yonma" if max_val >= 0.90 else "sanma"

    return game_type, max_val


# =========================
# OCR処理
# =========================

for filename in os.listdir(IMAGE_DIR):

    if not filename.lower().endswith((".png", ".jpg", ".jpeg")):
        continue

    image_path = os.path.join(IMAGE_DIR, filename)

    image_hash = get_image_hash(image_path)

    if image_hash in registered_hashes:
        print("登録済み画像のためスキップ:", filename)
        continue

    print("画像:", image_path)

    img_cv = imread_unicode(image_path)

    if img_cv is None:
        print("画像読み込み失敗")
        continue

    h, w = img_cv.shape[:2]

    target_w = 1920
    scale = target_w / w
    target_h = int(h * scale)

    img_cv = cv2.resize(img_cv, (target_w, target_h))

    game_type, fourth_rank_match = detect_game_type(img_cv)
    player_count = 3 if game_type == "sanma" else 4
    game_type_label = "サンマ" if game_type == "sanma" else "四麻"

    print(
        "対局種別:",
        game_type_label,
        "4位一致率:",
        round(fourth_rank_match, 3)
    )

    for rank in range(1, player_count + 1):

        template_path = os.path.join(TEMPLATE_DIR, f"{rank}.png")
        template = cv2.imread(template_path)

        if template is None:
            print("テンプレート読み込み失敗:", template_path)
            continue

        result = cv2.matchTemplate(
            img_cv,
            template,
            cv2.TM_CCOEFF_NORMED
        )

        _, max_val, _, max_loc = cv2.minMaxLoc(result)

        x, y = max_loc

        print(rank, "一致率:", round(max_val, 3), "場所:", (x, y))

        if rank == 1:
            name_right = 440 if game_type == "yonma" else 560

            name_img = img_cv[
                max(y + 0, 0):min(y + 55, img_cv.shape[0]),
                max(x + 170, 0):min(x + name_right, img_cv.shape[1])
            ]

            score_img = img_cv[
                max(y + 40, 0):min(y + 140, img_cv.shape[0]),
                max(x + 200, 0):min(x + 460, img_cv.shape[1])
            ]

            point_img = img_cv[
                max(y + 45, 0):min(y + 125, img_cv.shape[0]),
                max(x + 430, 0):min(x + 580, img_cv.shape[1])
            ]

        else:
            name_img = img_cv[
                max(y - 8, 0):min(y + 42, img_cv.shape[0]),
                max(x + 165, 0):min(x + 380, img_cv.shape[1])
            ]

            score_img = img_cv[
                max(y + 35, 0):min(y + 100, img_cv.shape[0]),
                max(x + 160, 0):min(x + 380, img_cv.shape[1])
            ]

            point_img = img_cv[
                max(y + 40, 0):min(y + 85, img_cv.shape[0]),
                max(x + 370, 0):min(x + 470, img_cv.shape[1])
            ]

        imwrite_unicode(
            os.path.join(OUTPUT_DIR, f"name_rank{rank}_{filename}"),
            name_img
        )

        name_text = ocr_name_easy(
            name_img,
            f"name_easy_rank{rank}_{filename}"
        )

        name_manual_input = False

        if not name_text:
            manual_name = ask_name_manual_gui(
                filename,
                rank
            )

            if manual_name:
                name_text = manual_name
                name_manual_input = True

        score_big = cv2.resize(score_img, None, fx=4, fy=4)
        score_gray = cv2.cvtColor(score_big, cv2.COLOR_BGR2GRAY)

        _, score_bin = cv2.threshold(
            score_gray,
            200,
            255,
            cv2.THRESH_BINARY
        )

        score_bin = cv2.bitwise_not(score_bin)

        imwrite_unicode(
            os.path.join(OUTPUT_DIR, f"score_bin_rank{rank}_{filename}"),
            score_bin
        )

        minus_detected = detect_minus_from_score_bin(score_bin)

        score_text = pytesseract.image_to_string(
            Image.fromarray(score_bin),
            lang="eng",
            config="--psm 7 -c tessedit_char_whitelist=0123456789"
        ).strip()

        score_text = score_text.replace(" ", "")

        score_numbers = re.findall(r"\d+", score_text)

        score_value = score_numbers[0] if score_numbers else ""

        if score_value and minus_detected:
            score_value = "-" + score_value

        point_big = cv2.resize(point_img, None, fx=5, fy=5)
        hsv = cv2.cvtColor(point_big, cv2.COLOR_BGR2HSV)

        lower_green = (35, 40, 40)
        upper_green = (95, 255, 255)

        green_mask = cv2.inRange(
            hsv,
            lower_green,
            upper_green
        )

        lower_red1 = (0, 40, 40)
        upper_red1 = (15, 255, 255)

        lower_red2 = (160, 40, 40)
        upper_red2 = (180, 255, 255)

        red_mask1 = cv2.inRange(
            hsv,
            lower_red1,
            upper_red1
        )

        red_mask2 = cv2.inRange(
            hsv,
            lower_red2,
            upper_red2
        )

        red_mask = cv2.bitwise_or(red_mask1, red_mask2)

        green_text, green_num = ocr_digits_from_mask(green_mask)
        red_text, red_num = ocr_digits_from_mask(red_mask)

        point_value = ""

        if green_num:
            point_value = "+" + green_num

        elif red_num:
            point_value = "-" + red_num

        check_flag = "OK"
        reasons = []

        if not name_text:
            check_flag = "要確認"
            reasons.append("名前OCR失敗")

        if name_manual_input:
            reasons.append("名前手入力")

        if not score_value:
            check_flag = "要確認"
            reasons.append("素点OCR失敗")

        if not point_value:
            check_flag = "要確認"
            reasons.append("順位点OCR失敗")

        if str(score_value).startswith("-"):
            check_flag = "要確認"
            reasons.append("素点マイナス確認")

        rows.append({
            "画像ハッシュ": image_hash,
            "画像ファイル": filename,
            "対局種別": game_type_label,
            "順位": rank,
            "名前OCR": name_text,
            "素点OCR": score_text,
            "素点": score_value,
            "順位点": point_value,
            "確認フラグ": check_flag,
            "確認理由": " / ".join(reasons),
            "一致率": round(max_val, 3),
            "検出X": x,
            "検出Y": y
        })

        print("順位:", rank)
        print("名前OCR:", name_text)
        print("素点:", score_value)
        print("順位点:", point_value)
        print("確認:", check_flag)

    registered_hashes.add(image_hash)

    print("完了")


# =========================
# Excel出力
# =========================

def generate_reports(game_type):
    output_dir = os.path.join(OUTPUT_DIR, game_type)
    os.makedirs(output_dir, exist_ok=True)

    player_count = 3 if game_type == "sanma" else 4
    last_rank = player_count
    game_type_label = "\u30b5\u30f3\u30de" if game_type == "sanma" else "\u56db\u9ebb"
    master_path = os.path.join(output_dir, "master_result.xlsx")
    old_df = old_dfs[game_type]

    new_df = pd.DataFrame([
        row for row in rows
        if row.get("\u5bfe\u5c40\u7a2e\u5225") == game_type_label
    ])

    if not old_df.empty:
        all_df = pd.concat(
            [old_df, new_df],
            ignore_index=True
        )
    else:
        all_df = new_df.copy()


    if all_df.empty:
        print("新規データなし")

    else:
        for col in [
            "画像ハッシュ",
            "画像ファイル",
            "名前OCR",
            "素点",
            "順位点"
        ]:
            if col in all_df.columns:
                all_df[col] = all_df[col].astype(str).str.strip()

        all_df = all_df.drop_duplicates(
            subset=[
                "画像ハッシュ",
                "順位"
            ],
            keep="first"
        )

        # =========================
        # 名前OCR確認・一括補正
        # =========================

        alias_dict = load_name_aliases()

        if alias_dict:
            all_df["名前OCR"] = (
                all_df["名前OCR"]
                .replace(alias_dict)
            )

        confirmation_df = new_df.copy()

        if not confirmation_df.empty:
            for col in ["画像ファイル", "名前OCR", "順位"]:
                if col in confirmation_df.columns:
                    confirmation_df[col] = (
                        confirmation_df[col]
                        .astype(str)
                        .str.strip()
                    )

            if alias_dict:
                confirmation_df["名前OCR"] = (
                    confirmation_df["名前OCR"]
                    .replace(alias_dict)
                )

            existing_player_names = set()

            if not old_df.empty and "名前OCR" in old_df.columns:
                existing_player_names = set(
                    old_df["名前OCR"]
                    .dropna()
                    .astype(str)
                    .str.strip()
                    .replace(alias_dict)
                )

            confirmation_df = confirmation_df[
                ~confirmation_df["名前OCR"].isin(existing_player_names)
            ]

        name_sources = {}

        if "名前OCR" in confirmation_df.columns:
            current_player_names = sorted(
                confirmation_df["名前OCR"]
                .dropna()
                .astype(str)
                .unique()
            )
            source_groups = confirmation_df.groupby("名前OCR")
        else:
            current_player_names = []
            source_groups = []

        for name, group in source_groups:
            clean_name = clean_name_text(name)

            if not clean_name:
                continue

            sources = []

            for _, row in group.iterrows():
                filename = str(row.get("画像ファイル", "")).strip()
                rank = str(row.get("順位", "")).strip()

                if filename and filename != "nan":
                    source_label = filename

                    if rank and rank != "nan":
                        source_label += f" ({rank}位)"

                    try:
                        rank_for_path = str(int(float(rank)))
                    except Exception:
                        rank_for_path = rank

                    sources.append({
                        "label": source_label,
                        "preview_path": os.path.join(
                            OUTPUT_DIR,
                            f"name_easy_rank{rank_for_path}_{filename}"
                        )
                    })

            name_sources[clean_name] = sorted(
                sources,
                key=lambda source: (
                    not os.path.exists(source["preview_path"]),
                    source["label"]
                )
            )

        correction_dict = confirm_names_gui(
            current_player_names,
            name_sources
        )

        if correction_dict:
            all_df["名前OCR"] = (
                all_df["名前OCR"]
                .replace(correction_dict)
            )

            alias_dict.update(correction_dict)
            save_name_aliases(alias_dict)

            if "確認理由" in all_df.columns:
                corrected_names = set(correction_dict.values())
                all_df.loc[
                    all_df["名前OCR"].isin(corrected_names),
                    "確認理由"
                ] = (
                    all_df.loc[
                        all_df["名前OCR"].isin(corrected_names),
                        "確認理由"
                    ].astype(str).replace("nan", "") + " / 名前一覧で補正"
                ).str.strip(" /")

        all_df["素点_数値"] = pd.to_numeric(
            all_df["素点"],
            errors="coerce"
        )

        all_df["順位点_数値"] = pd.to_numeric(
            all_df["順位点"].astype(str).str.replace("+", "", regex=False),
            errors="coerce"
        )

        all_df["順位_数値"] = pd.to_numeric(
            all_df["順位"],
            errors="coerce"
        )

        with pd.ExcelWriter(
            master_path,
            engine="openpyxl"
        ) as writer:
            all_df.to_excel(
                writer,
                index=False,
                sheet_name="master"
            )

        print("master_result.xlsx 更新完了:", master_path)

        summary_df = (
            all_df
            .pivot_table(
                index="名前OCR",
                columns="順位_数値",
                values="画像ファイル",
                aggfunc="count",
                fill_value=0
            )
            .reset_index()
        )

        summary_df.columns.name = None

        summary_df = summary_df.rename(
            columns={
                1: "1位回数",
                2: "2位回数",
                3: "3位回数",
                4: "4位回数"
            }
        )

        for col in [
            "1位回数",
            "2位回数",
            "3位回数",
            "4位回数"
        ]:
            if col not in summary_df.columns:
                summary_df[col] = 0

        rank_count_columns = [
            f"{rank}位回数"
            for rank in range(1, player_count + 1)
        ]

        summary_df["対局数"] = summary_df[rank_count_columns].sum(axis=1)

        summary_df["トップ率"] = (
            summary_df["1位回数"]
            / summary_df["対局数"]
            * 100
        ).round(1)

        summary_df["ラス率"] = (
            summary_df[f"{last_rank}位回数"]
            / summary_df["対局数"]
            * 100
        ).round(1)

        summary_df["平均順位"] = (
            (
                summary_df["1位回数"] * 1
                + summary_df["2位回数"] * 2
                + summary_df["3位回数"] * 3
                + (
                    summary_df["4位回数"] * 4
                    if player_count == 4
                    else 0
                )
            )
            / summary_df["対局数"]
        ).round(2)

        score_sum = (
            all_df
            .groupby("名前OCR")["素点_数値"]
            .sum()
            .reset_index(name="素点合計")
        )

        point_sum = (
            all_df
            .groupby("名前OCR")["順位点_数値"]
            .sum()
            .reset_index(name="順位点合計")
        )

        fly_count = (
            all_df[all_df["素点_数値"] < 0]
            .groupby("名前OCR")
            .size()
            .reset_index(name="飛び回数")
        )

        summary_df = summary_df.merge(score_sum, on="名前OCR", how="left")
        summary_df = summary_df.merge(point_sum, on="名前OCR", how="left")
        summary_df = summary_df.merge(fly_count, on="名前OCR", how="left")

        summary_df["飛び回数"] = summary_df["飛び回数"].fillna(0).astype(int)

        summary_df["飛び率"] = (
            summary_df["飛び回数"]
            / summary_df["対局数"]
            * 100
        ).round(1)

        summary_df = summary_df[
            [
                "名前OCR",
                "対局数",
                *rank_count_columns,
                "トップ率",
                "ラス率",
                "飛び回数",
                "飛び率",
                "平均順位",
                "素点合計",
                "順位点合計"
            ]
        ]

        vs_rows = []

        for game_hash, game_df in all_df.groupby("画像ハッシュ"):

            game_df = game_df.copy()

            game_df = game_df.dropna(
                subset=[
                    "名前OCR",
                    "順位_数値"
                ]
            )

            if len(game_df) < 2:
                continue

            players = game_df.to_dict("records")

            for player in players:
                player_name = player["名前OCR"]
                player_rank = player["順位_数値"]

                for opponent in players:
                    opponent_name = opponent["名前OCR"]
                    opponent_rank = opponent["順位_数値"]

                    if player_name == opponent_name:
                        continue

                    win_flag = 1 if player_rank < opponent_rank else 0
                    lose_flag = 1 if player_rank > opponent_rank else 0

                    vs_rows.append({
                        "名前OCR": player_name,
                        "対戦相手": opponent_name,
                        "画像ハッシュ": game_hash,
                        "自分順位": player_rank,
                        "相手順位": opponent_rank,
                        "勝ち": win_flag,
                        "負け": lose_flag
                    })

        vs_raw_df = pd.DataFrame(vs_rows)

        if not vs_raw_df.empty:

            vs_summary_df = (
                vs_raw_df
                .groupby(
                    [
                        "名前OCR",
                        "対戦相手"
                    ]
                )
                .agg(
                    同卓回数=("画像ハッシュ", "count"),
                    勝ち数=("勝ち", "sum"),
                    負け数=("負け", "sum")
                )
                .reset_index()
            )

            vs_summary_df["勝率"] = (
                vs_summary_df["勝ち数"]
                / vs_summary_df["同卓回数"]
                * 100
            ).round(1)

            vs_summary_df = vs_summary_df[
                [
                    "名前OCR",
                    "対戦相手",
                    "同卓回数",
                    "勝ち数",
                    "負け数",
                    "勝率"
                ]
            ]

        else:
            vs_summary_df = pd.DataFrame(
                columns=[
                    "名前OCR",
                    "対戦相手",
                    "同卓回数",
                    "勝ち数",
                    "負け数",
                    "勝率"
                ]
            )

        # =========================
        # summary_result.xlsx 作成
        # =========================

        summary_path = os.path.join(
            output_dir,
            "summary_result.xlsx"
        )

        player_names = sorted(
            summary_df["名前OCR"]
            .dropna()
            .astype(str)
            .unique()
        )

        target_players = select_players_gui(
            player_names
        )

        with pd.ExcelWriter(
            summary_path,
            engine="openpyxl"
        ) as writer:

            summary_df.to_excel(
                writer,
                index=False,
                sheet_name="summary"
            )

            vs_summary_df.to_excel(
                writer,
                index=False,
                sheet_name="vs_summary"
            )

            vs_raw_df.to_excel(
                writer,
                index=False,
                sheet_name="vs_raw"
            )

            # 選択したプレイヤーだけ個別シートを作成
            for player_name in target_players:

                safe_name = str(player_name)

                for ng in ["\\", "/", "*", "?", ":", "[", "]"]:
                    safe_name = safe_name.replace(ng, "")

                sheet_name = "P_" + safe_name

                if len(sheet_name) > 31:
                    sheet_name = sheet_name[:31]

                player_vs_df = vs_summary_df[
                    vs_summary_df["名前OCR"] == player_name
                ].copy()

                player_vs_df = player_vs_df[
                    [
                        "対戦相手",
                        "同卓回数",
                        "勝ち数",
                        "負け数",
                        "勝率"
                    ]
                ]

                player_vs_df = player_vs_df.sort_values(
                    by=[
                        "同卓回数",
                        "勝率"
                    ],
                    ascending=[
                        False,
                        False
                    ]
                )

                player_vs_df.to_excel(
                    writer,
                    index=False,
                    sheet_name=sheet_name
                )

                # =========================
                # 選択プレイヤーの戦績ページを作成
                # =========================

                stat_sheet_name = "戦績_" + safe_name

                if len(stat_sheet_name) > 31:
                    stat_sheet_name = stat_sheet_name[:31]

                player_summary_df = summary_df[
                    summary_df["名前OCR"] == player_name
                ].copy()

                player_detail_df = all_df[
                    all_df["名前OCR"] == player_name
                ].copy()

                if not player_summary_df.empty:
                    summary_row = player_summary_df.iloc[0]

                    stat_rows = [
                        ["対局数", summary_row.get("対局数", "")],
                        ["平均順位", summary_row.get("平均順位", "")],
                        ["トップ率", summary_row.get("トップ率", "")],
                        ["ラス率", summary_row.get("ラス率", "")],
                        ["飛び回数", summary_row.get("飛び回数", "")],
                        ["飛び率", summary_row.get("飛び率", "")],
                        ["素点合計", summary_row.get("素点合計", "")],
                        ["順位点合計", summary_row.get("順位点合計", "")],
                    ]
                else:
                    stat_rows = []

                stat_df = pd.DataFrame(
                    stat_rows,
                    columns=["項目", "値"]
                )

                rank_rows = []

                if not player_summary_df.empty:
                    summary_row = player_summary_df.iloc[0]

                    rank_rows = [
                        [f"{rank}位", summary_row.get(f"{rank}位回数", 0)]
                        for rank in range(1, player_count + 1)
                    ]

                rank_df = pd.DataFrame(
                    rank_rows,
                    columns=["順位", "回数"]
                )

                player_vs_page_df = player_vs_df.copy()

                player_game_df = player_detail_df[
                    [
                        "画像ファイル",
                        "順位",
                        "素点",
                        "順位点",
                        "確認フラグ",
                        "確認理由"
                    ]
                ].copy()

                player_game_df = player_game_df.sort_values(
                    by=["画像ファイル"],
                    ascending=True
                )

                stat_df.to_excel(
                    writer,
                    index=False,
                    sheet_name=stat_sheet_name,
                    startrow=2
                )

                rank_df.to_excel(
                    writer,
                    index=False,
                    sheet_name=stat_sheet_name,
                    startrow=2,
                    startcol=3
                )

                worksheet = writer.sheets[stat_sheet_name]
                worksheet.cell(row=1, column=1, value=f"{player_name}の戦績")

                player_vs_page_df.to_excel(
                    writer,
                    index=False,
                    sheet_name=stat_sheet_name,
                    startrow=13
                )

                player_game_df.to_excel(
                    writer,
                    index=False,
                    sheet_name=stat_sheet_name,
                    startrow=13,
                    startcol=6
                )

            apply_excel_styles(writer.book)

        print("summary_result.xlsx 作成完了:", summary_path)


for game_type in ["yonma", "sanma"]:
    generate_reports(game_type)


# =========================
# 一時画像削除
# =========================

for file in os.listdir(OUTPUT_DIR):
    if file.lower().endswith((".png", ".jpg", ".jpeg")):
        try:
            os.remove(os.path.join(OUTPUT_DIR, file))
        except Exception:
            pass

print("一時画像を削除しました")








